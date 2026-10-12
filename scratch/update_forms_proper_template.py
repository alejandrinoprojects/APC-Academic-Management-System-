import os, subprocess
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FOLDER = r"C:\Users\Dean Alejandrino\Downloads\New folder (3)"
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
USER_DIR = r"C:\Users\Dean Alejandrino\.gemini\antigravity\scratch\edge_profile_forms"
os.makedirs(USER_DIR, exist_ok=True)

# SIMPLIFIED, EASY-TO-UNDERSTAND ACTIONS TAKEN
EMBILE_ACTIONS = {
    # Table 1: Comments
    1: {
        1: "Checked the project scope from Chapter 1 against Chapter 3 and the actual software. All required features—such as curriculum editing, flowchart with prerequisites, spreadsheet, OBE mapping, user roles, task delegations, and reports—are now fully included in the system."
    },
    # Table 2: Revisions (none)
    2: {},
    # Table 3: Additional Requirements
    3: {
        1: "Added admin settings and maintenance tools to the system. This allows authorized users to manage user roles, update institutional statements (Mission, Vision, GA, PEO, and SO), and view the system audit log.",
        2: "Updated all official document and report layouts to automatically show the report title, creation date, print date, and the name and role of the user who printed it."
    }
}

PERUDA_ACTIONS = {
    # Table 1: Comments
    1: {
        1: "Added a clear warning badge (yellow warning icon) and a pop-up notice in the curriculum map when a capstone or final course does not reach the Demonstrative (D) level.",
        2: "Updated the summary to display the total number of courses mapped to each Student Outcome (course count) instead of calculating credit units per SO.",
        3: "Explained and illustrated the prerequisite cycle check in Chapter 3 (Figure 3.6). In the flowchart editor, if a user accidentally makes two courses prerequisites of each other (a circular loop), the system displays a red error banner showing the loop and prevents saving until fixed.",
        4: "Clearly placed the flowchart feature under the Curriculum Management subsystem in Chapter 3 (Section 3.1 & Figure 3.1), showing it as a core tool for curriculum design and exporting images.",
        5: "Cross-checked every project objective in Chapter 1 against the project scope items to make sure each objective has a corresponding feature in the system."
    },
    # Table 2: Revisions
    2: {
        1: "Separated curriculum handling into three distinct areas: (1) a creation tool to start new curricula, (2) an editing tool for working drafts (like CPE 2027), and (3) a read-only viewer for approved/published curricula (like CPE 2026) to prevent accidental changes.",
        2: "Changed the labels to show the curriculum version (e.g., CPE 2026) rather than student year levels (e.g., 1st Year to 4th Year), since students advance every year while the curriculum stays the same.",
        3: "Standardized batch names across the entire system to include both curriculum version and batch year (for example, 'CPE 2026 [Batch 2026]').",
        4: "Removed confusing extra parentheses from menus, dropdowns, and diagrams. Replaced them with simple academic year labels (for example, 'AY 2026-2027').",
        5: "Clarified this in Chapter 3 using the example: if Capstone 1 is Introductory (I), Capstone 2 should not skip directly to Demonstrative (D) without an Enabling (E) course in between.",
        6: "Made learning progression checks advisory warnings only. The system shows an amber warning notice but still allows the user to save changes if needed.",
        7: "Changed the text from 'Credit Units per SO' to 'Courses per Student Outcome' to accurately describe the course count shown.",
        8: "Removed unapproved extra features from the chapter and system, keeping only features that were approved in Chapter 1.",
        9: "Simplified the wording across Chapter 3, replacing overly complex terms with plain, standard words so everything is easy to understand."
    },
    # Table 3: Additional Requirements
    3: {
        1: "Added a temporary review screen when importing Excel files. It checks all rows for duplicate codes and missing info before anything is saved into the database.",
        2: "Updated the import window to show the exact row and column where any error occurred (for example, 'Row 12, Column C: Missing Units') so the user can fix it easily.",
        3: "Blocked Excel files that contain merged cells from being imported, and added an error message reminding the user that rows must be flat and unmerged.",
        4: "Created a dedicated page to add and edit the Vision, Mission, and Core Values with version numbers, so updating them does not change older curricula.",
        5: "Added dedicated tools to create, edit, and update Graduate Attributes, PEOs, and Student Outcomes as distinct items with their own records and versions.",
        6: "Gave each item its own version number and linked them hierarchically: each SO version points to its parent PEO version, and each PEO points to its parent GA version.",
        7: "Locked all published curricula (like CPE 2026, 2025, 2024) to read-only mode. Any changes require creating a new draft revision."
    }
}

YONG_ACTIONS = {
    # Table 1: Comments
    1: {
        1: "Added a flexible template system so that when CHED issues a new format, a new document template can be created without affecting old, existing document formats.",
        2: "Simplified faculty delegation by removing rigid cluster requirements. Program Directors can now assign specific courses directly to faculty members."
    },
    # Table 2: Revisions
    2: {
        1: "Froze the course code and course title columns in place so they stay visible when scrolling horizontally through the spreadsheet.",
        2: "Fixed the comparative summary table: corrected the column headers, removed unnecessary rows, and ensured unit counts match official registrar standards.",
        3: "Removed the full-time and part-time status selection from the delegation form so it only asks for the faculty member's name and assigned course."
    },
    # Table 3: Additional Requirements
    3: {
        1: "Moved institutional statements out of the code and into an editable settings tool where each update is saved as a new version.",
        2: "Added a summary box at the bottom of each term column showing the total lecture hours, lab hours, and credit units.",
        3: "Added a tool that lets curriculum chairs create, edit, and delete course groups (such as General Education, Core, or Electives) instead of using a fixed list.",
        4: "Connected the course group dropdown in the spreadsheet to the custom course group tool so users can manage group options directly.",
        5: "Added editable CHED minimum requirements tables and a file parser that reads and organizes course lists from uploaded files into the comparative sheet.",
        6: "Added buttons on the comparative summary page that allow users to add new rows and columns as needed.",
        7: "Added a version dropdown on the comparative page so users can select and compare current drafts with any previous curriculum version."
    }
}

FORMS = [
    {
        "docx_name": "KapTau_-_Chapter_3_Approval_Form__EMBILE_ (1).docx",
        "pdf_name": "KapTau_Chapter_3_Approval_Form_EMBILE.pdf",
        "panel_member": "ENGR. ARNETT EMBILE",
        "actions_map": EMBILE_ACTIONS
    },
    {
        "docx_name": "KapTau_-_Chapter_3_Approval_Form__PERUDA_ (1).docx",
        "pdf_name": "KapTau_Chapter_3_Approval_Form_PERUDA.pdf",
        "panel_member": "SERGIO R. PERUDA JR.",
        "actions_map": PERUDA_ACTIONS
    },
    {
        "docx_name": "KapTau_-_Chapter_3_Approval_Form__YONG_ (1).docx",
        "pdf_name": "KapTau_Chapter_3_Approval_Form_YONG.pdf",
        "panel_member": "ENGR. CYNTHIA YONG",
        "actions_map": YONG_ACTIONS
    }
]

def update_docx(form_info):
    docx_path = os.path.join(FOLDER, form_info["docx_name"])
    doc = docx.Document(docx_path)
    
    # Table 0: Ensure Recommendation is EMPTY
    t0 = doc.tables[0]
    if len(t0.rows) > 1 and len(t0.rows[1].cells) > 1:
        t0.rows[1].cells[1].text = ""
        
    actions_map = form_info["actions_map"]
    for t_idx in [1, 2, 3]:
        if t_idx >= len(doc.tables):
            continue
        table = doc.tables[t_idx]
        t_actions = actions_map.get(t_idx, {})
        for r_idx in range(1, len(table.rows)):
            row = table.rows[r_idx]
            # Column 1: Actions Taken
            action_text = t_actions.get(r_idx, "")
            if len(row.cells) > 1:
                row.cells[1].text = ""
                p = row.cells[1].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if action_text:
                    r = p.add_run(action_text)
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
            # Column 2: Remarks - MUST BE COMPLETELY EMPTY
            if len(row.cells) > 2:
                row.cells[2].text = ""
                
    doc.save(docx_path)
    print(f"Updated DOCX: {form_info['docx_name']}")


def generate_proper_template_pdf(form_info):
    docx_path = os.path.join(FOLDER, form_info["docx_name"])
    doc = docx.Document(docx_path)
    
    # Extract comment texts from doc tables
    def get_table_rows(t_idx):
        if t_idx >= len(doc.tables):
            return []
        rows_data = []
        t = doc.tables[t_idx]
        for r_idx in range(1, len(t.rows)):
            row = t.rows[r_idx]
            c0 = row.cells[0].text.strip()
            c1 = row.cells[1].text.strip() if len(row.cells) > 1 else ""
            c2 = "" # Remarks strictly empty
            if c0:
                rows_data.append((c0, c1, c2))
        return rows_data

    t1_rows = get_table_rows(1)
    t2_rows = get_table_rows(2)
    t3_rows = get_table_rows(3)
    
    # Build clean HTML matching the exact institutional template
    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{form_info['panel_member']} - Approval Form</title>
<style>
  @page {{
    size: letter portrait;
    margin: 0.75in 0.85in;
    @bottom-center {{
      content: counter(page);
    }}
  }}
  body {{
    font-family: Arial, "Helvetica Neue", Helvetica, sans-serif;
    font-size: 10pt;
    line-height: 1.35;
    color: #000000;
    margin: 0;
    padding: 0;
  }}
  .header-top {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    font-size: 10pt;
    margin-bottom: 12pt;
  }}
  .school-title {{
    font-size: 10pt;
    font-weight: normal;
  }}
  .userinx-title {{
    font-size: 10pt;
    font-weight: bold;
    text-align: right;
  }}
  .form-title {{
    text-align: center;
    font-weight: bold;
    font-size: 11pt;
    margin: 0 0 2pt 0;
  }}
  .form-subtitle {{
    text-align: center;
    font-size: 9.5pt;
    font-style: italic;
    margin: 0 0 14pt 0;
  }}
  .greeting {{
    text-align: center;
    font-weight: bold;
    margin-bottom: 8pt;
  }}
  p.body-text {{
    text-align: justify;
    margin: 0 0 6pt 0;
    font-size: 9.5pt;
    line-height: 1.35;
  }}
  .rec-heading {{
    font-weight: bold;
    margin-top: 8pt;
    margin-bottom: 4pt;
    font-size: 9.5pt;
  }}
  .rec-list {{
    margin: 0 0 10pt 0;
    padding-left: 18pt;
    font-size: 9pt;
    line-height: 1.35;
  }}
  .rec-list li {{
    margin-bottom: 3pt;
  }}
  .rec-list strong {{
    font-weight: bold;
  }}
  table.form-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 8pt 0 14pt 0;
    font-size: 9pt;
  }}
  table.form-table th, table.form-table td {{
    border: 1px solid #000;
    padding: 5pt 7pt;
    vertical-align: top;
    line-height: 1.3;
  }}
  table.form-table th {{
    background-color: #f2f2f2;
    font-weight: bold;
    text-align: left;
  }}
  .section-label {{
    font-weight: bold;
    font-size: 10pt;
    margin-top: 12pt;
    margin-bottom: 4pt;
  }}
  .signature-block {{
    margin-top: 24pt;
    font-size: 9.5pt;
    line-height: 2.0;
    page-break-inside: avoid;
  }}
  .sig-line {{
    display: flex;
    justify-content: space-between;
  }}
</style>
</head>
<body>

  <div class="header-top">
    <div class="school-title">School of Engineering</div>
    <div class="userinx-title">USERINX</div>
  </div>

  <div class="form-title">User Interface / User Experience Proposed Software Design Project Title Pre-approval Form</div>
  <div class="form-subtitle">(Refer to attached initial title proposal)</div>

  <div class="greeting">Dear Panel Member,</div>

  <p class="body-text">
    Attached is the Chapter 1 and 2 of the title proposal for USERINX CPE251 containing Chapter 1: Background of the Project, Problem Identification, Statement of the Problem, Project Objectives, Scope of Work, Assumptions, Constrains, Significance of the Project, and Definition of Terms; Chapter 2: Stakeholder Analysis, Current Business Process, User Analysis, User Journey Map, Pain Point Analysis, Gap Analysis, Review of Related Literatures and Requirement Analysis.
  </p>
  <p class="body-text">
    Please determine if the attached title proposal/s is approved for title defense in terms of feasibility and relevance to CpE's field of specialization.
  </p>
  <p class="body-text">
    Once approved, the proponents are expected to proceed to project proposal documentation (Chapters 1 to 3) for project proposal defense.
  </p>

  <div class="rec-heading">Possible Recommendations:</div>
  <ul class="rec-list">
    <li><strong>Approved</strong> – The proposed title is accepted without revisions. The project team may proceed to the Proposal Defense.</li>
    <li><strong>Approved with Revisions</strong> – The proposed title is accepted subject to the completion of the required revisions. The project team may proceed to the Proposal Defense upon compliance with the panel's recommendations.</li>
    <li><strong>Change Topic</strong> – The proposed title is not approved. The project team shall propose a new project topic that addresses the recommendations of the evaluation panel.</li>
  </ul>
  <p class="body-text">
    If the project receives a Change Topic result during the scheduled Title Defense, the remaining defense schedule allocated for the UI/UX course shall be used for the Proposal Defense of the revised project. Any additional proposal defense, if required, shall be conducted during the next offering of the UI/UX course.
  </p>

  <!-- TABLE 0 -->
  <table class="form-table" style="margin-bottom: 12pt;">
    <thead>
      <tr>
        <th style="width: 72%;">PROPOSED TITLE</th>
        <th style="width: 28%; text-align: center;">RECOMMENDATION</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td style="font-weight: bold; padding: 8pt 7pt;">Development of a Centralized Curriculum Management and Mapping System with Role-Based Access Control and Audit Logging</td>
        <td></td>
      </tr>
    </tbody>
  </table>

  <!-- TABLE 1 -->
  <div class="section-label">Comments</div>
  <table class="form-table">
    <thead>
      <tr>
        <th style="width: 44%;">Comments</th>
        <th style="width: 44%;">Actions Taken</th>
        <th style="width: 12%; text-align: center;">Remarks</th>
      </tr>
    </thead>
    <tbody>
"""
    for c0, c1, c2 in t1_rows:
        html += f"""      <tr>
        <td>{c0.replace(chr(10), '<br>')}</td>
        <td>{c1}</td>
        <td></td>
      </tr>\n"""
        
    html += """    </tbody>
  </table>
"""

    if t2_rows:
        html += """  <!-- TABLE 2 -->
  <div class="section-label">For revisions</div>
  <table class="form-table">
    <thead>
      <tr>
        <th style="width: 44%;">Revisions</th>
        <th style="width: 44%;">Actions Taken</th>
        <th style="width: 12%; text-align: center;">Remarks</th>
      </tr>
    </thead>
    <tbody>
"""
        for c0, c1, c2 in t2_rows:
            html += f"""      <tr>
        <td>{c0.replace(chr(10), '<br>')}</td>
        <td>{c1}</td>
        <td></td>
      </tr>\n"""
        html += """    </tbody>
  </table>
"""

    if t3_rows:
        html += """  <!-- TABLE 3 -->
  <div class="section-label">Additional requirements (additional objectives that needs to be included in the complete proposal)</div>
  <table class="form-table">
    <thead>
      <tr>
        <th style="width: 44%;">Additional Requirements</th>
        <th style="width: 44%;">Actions Taken</th>
        <th style="width: 12%; text-align: center;">Remarks</th>
      </tr>
    </thead>
    <tbody>
"""
        for c0, c1, c2 in t3_rows:
            html += f"""      <tr>
        <td>{c0.replace(chr(10), '<br>')}</td>
        <td>{c1}</td>
        <td></td>
      </tr>\n"""
        html += """    </tbody>
  </table>
"""

    html += f"""  <div class="signature-block">
    <div><strong>NAME OF PANEL MEMBER:</strong> {form_info['panel_member']}&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<strong>SIGNATURE :</strong> _________________________________</div>
    <div><strong>Date Signed :</strong> _________________________________</div>
  </div>

</body>
</html>
"""

    html_path = os.path.join(FOLDER, form_info["pdf_name"].replace(".pdf", ".html"))
    pdf_path = os.path.join(FOLDER, form_info["pdf_name"])
    
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
        
    url = "file:///" + os.path.abspath(html_path).replace("\\", "/")
    cmd = [
        EDGE_PATH,
        "--headless",
        "--disable-gpu",
        f"--user-data-dir={USER_DIR}",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        url
    ]
    subprocess.run(cmd, capture_output=True, text=True)
    if os.path.exists(html_path):
        os.remove(html_path)
    print(f"Generated clean PDF: {form_info['pdf_name']} (size: {os.path.getsize(pdf_path)})")


if __name__ == "__main__":
    for f_info in FORMS:
        update_docx(f_info)
        generate_proper_template_pdf(f_info)
    print("All forms updated and converted successfully.")
