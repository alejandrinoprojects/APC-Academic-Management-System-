import docx
import re

doc = docx.Document(r"C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Master_Cleaned.docx")
print("Total paragraphs:", len(doc.paragraphs))
for i in range(221, 295):
    p = doc.paragraphs[i]
    xml = p._p.xml
    imgs = re.findall(r'r:embed="(rId\d+)"', xml)
    if imgs or p.text.strip():
        rels = [doc.part.rels[r].target_ref for r in imgs] if imgs else []
        print(f"P[{i}]: text='{p.text.strip()[:50]}' imgs={rels}")
