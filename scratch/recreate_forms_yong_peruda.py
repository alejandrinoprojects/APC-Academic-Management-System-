import os, subprocess, shutil
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC_DIR = r"C:\Users\Dean Alejandrino\Downloads"
OUT_DIR = r"C:\Users\Dean Alejandrino\Documents\antigravity\jolly-davinci\recreated_approval_forms"
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
USER_DIR = r"C:\Users\Dean Alejandrino\.gemini\antigravity\scratch\edge_profile_recreated_forms"

os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(USER_DIR, exist_ok=True)

PERUDA_ACTIONS = {
    # Table 1: Comments
    1: {
        1: "Added a yellow warning bar and an 'Advisory' button at the top of the curriculum matrix and flowchart. When a final course (like Capstone) is missing its Demonstrative (D) mark, the yellow button appears. Clicking it opens a popup that explains what is missing, but it still lets the user save their work.",
        2: "Removed the credit unit and contact hour calculations. The system now shows the simple count of how many courses are linked to each Student Outcome (SO a to m).",
        3: "Added loop detection in the Flowchart editor. If a user accidentally makes two courses prerequisites of each other (like Course A requires B, and B requires A), the system shows a red warning banner and stops the user from saving the loop.",
        4: "Placed the flowchart, course line highlighting, and image export directly under Curriculum Management as part of the main curriculum builder.",
        5: "Reviewed all project objectives to make sure every required feature is fully supported by the system."
    },
    # Table 2: Revisions
    2: {
        1: "Separated curriculum actions into three clear areas: (1) creating a new curriculum, (2) editing an active draft (like AY 2027), and (3) viewing published curricula in read-only mode (like AY 2026) where editing is locked.",
        2: "Updated the screen headers, dropdowns, and labels to show the curriculum version (like 'CPE 2026' or 'CPE 2027') instead of student year levels (1st Year to 4th Year).",
        3: "Labeled batches using both the curriculum version and entering batch year (for example, 'CPE 2026 [Batch 2026]') on the Curriculum Workbench, Directory, and Version History screens.",
        4: "Cleaned up all dropdowns and menus to remove confusing parentheses, using simple labels like 'AY 2026-2027' and 'CPE 2026'.",
        5: "Added an example to clarify progression warnings: if Course 1 is Introductory (I), Course 2 cannot jump straight to Demonstrative (D) without an Enabling (E) course in between. The system flags this skipped step with a yellow warning tag.",
        6: "Made sure all progression warnings are advisory only. They show up as yellow badges in the table and in a popup window to alert the user, but the user can still save their changes.",
        7: "Updated the Student Outcomes matrix to clearly display the count of courses mapped to each outcome rather than credit units.",
        8: "Removed unapproved extra features from the system, keeping only the approved features.",
        9: "Simplified all wording across the system and documentation, replacing confusing jargon with clear, standard explanations."
    },
    # Table 3: Additional Requirements
    3: {
        1: "Added a preview screen when uploading Excel files. The system checks for duplicate course codes and missing fields before saving anything, so incomplete or corrupted data is not saved.",
        2: "Added an error list right on the import screen that points out the exact row, column, and issue (such as a missing unit or duplicate code) before the file is imported.",
        3: "Set a rule that imported Excel files cannot have merged cells. If a file has merged cells, the system shows an error message asking the user to unmerge them first.",
        4: "Added an admin page to manage the University Mission, Vision, and Core Values. Each update is saved as a new version, so older curriculum records keep their original text.",
        5: "Added management screens to create, edit, and track versions for Graduate Attributes (GA), Program Educational Objectives (PEO), and Student Outcomes (SO).",
        6: "Linked the version history together: every SO version points to its base PEO version, and every PEO version points to its base GA version.",
        7: "Locked published curricula (like CPE 2026) so they cannot be edited. To make changes, the user must create a new draft version (like CPE 2027)."
    }
}

YONG_ACTIONS = {
    # Table 1: Comments
    1: {
        1: "Built the document generator using flexible templates. When CHED releases a new memo or format, administrators can add a new template without changing or breaking older documents.",
        2: "Simplified task delegation so Program Directors can assign courses or syllabus tasks directly to faculty members, with an optional grouping feature if the department chair chooses to use it."
    },
    # Table 2: Revisions
    2: {
        1: "Froze the Course Code and Course Title columns in the curriculum table so they stay visible on screen while scrolling sideways.",
        2: "Fixed the Comparative Summary table by correcting column headers, removing extra rows, and fixing the unit totals to match registrar standards.",
        3: "Removed the faculty rank and employment status fields from the task assignment screen, showing only the faculty member's name and assigned course or task."
    },
    # Table 3: Additional Requirements
    3: {
        1: "Moved the Vision, Mission, and Core Values into an admin settings page with version history, so they can be edited anytime without being hardcoded into the software.",
        2: "Added a summary row at the bottom of each term in the flowchart, showing the total lecture hours, lab hours, and credit units.",
        3: "Added a Course Group Manager that lets department chairs add, edit, and delete course categories (such as General Education, Core Courses, and Electives).",
        4: "Linked the course category dropdown in the curriculum table directly to the Course Group Manager, so any newly added group appears in the dropdown immediately.",
        5: "Added an editable form and an Excel import tool for CHED minimum requirements, allowing the system to read and fill in comparative summary tables automatically.",
        6: "Added buttons to add rows and columns directly inside the comparative summary table so users can add custom comparison metrics.",
        7: "Added a version selector on the comparative summary page so users can compare a new draft curriculum side-by-side with previous published versions."
    }
}

FORMS = [
    {
        "src_name": "KapTau_-_Chapter_3_Approval_Form__PERUDA_ (1).docx",
        "out_docx_name": "KapTau_-_Chapter_3_Approval_Form__PERUDA__RECREATED.docx",
        "out_pdf_name": "KapTau_-_Chapter_3_Approval_Form__PERUDA__RECREATED.pdf",
        "panel_member": "SERGIO R. PERUDA JR.",
        "actions_map": PERUDA_ACTIONS
    },
    {
        "src_name": "KapTau_-_Chapter_3_Approval_Form__YONG_ (1).docx",
        "out_docx_name": "KapTau_-_Chapter_3_Approval_Form__YONG__RECREATED.docx",
        "out_pdf_name": "KapTau_-_Chapter_3_Approval_Form__YONG__RECREATED.pdf",
        "panel_member": "EINSTEIN D. YONG",
        "actions_map": YONG_ACTIONS
    }
]

def recreate_docx(form_info):
    src_path = os.path.join(SRC_DIR, form_info["src_name"])
    out_docx_path = os.path.join(OUT_DIR, form_info["out_docx_name"])
    
    doc = docx.Document(src_path)
    
    # Table 0: Proposed Title & Recommendation (Ensure Recommendation is EMPTY)
    t0 = doc.tables[0]
    if len(t0.rows) > 1 and len(t0.rows[1].cells) > 1:
        t0.rows[1].cells[1].text = ""
        
    actions_map = form_info["actions_map"]
    for t_idx in [1, 2, 3]:
        if t_idx >= len(doc.tables):
            continue
        table = doc.tables[t_idx]
        t_actions = actions_map.get(t_idx, {})
        for r_idx in range(1, len(table.rows)):
            row = table.rows[r_idx]
            action_text = t_actions.get(r_idx, "")
            
            # Column 1: Actions Taken
            if len(row.cells) > 1:
                row.cells[1].text = ""
                p = row.cells[1].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if action_text:
                    r = p.add_run(action_text)
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    
            # Column 2: Remarks (MUST BE COMPLETELY EMPTY for panelist signing)
            if len(row.cells) > 2:
                row.cells[2].text = ""
                
    doc.save(out_docx_path)
    print(f"Created DOCX: {out_docx_path}")

def recreate_pdf(form_info):
    src_path = os.path.join(SRC_DIR, form_info["src_name"])
    doc = docx.Document(src_path)
    
    actions_map = form_info["actions_map"]
    
    def get_table_rows(t_idx):
        if t_idx >= len(doc.tables):
            return []
        rows_data = []
        t = doc.tables[t_idx]
        t_actions = actions_map.get(t_idx, {})
        for r_idx in range(1, len(t.rows)):
            row = t.rows[r_idx]
            c0 = row.cells[0].text.strip()
            c1 = t_actions.get(r_idx, "")
            if c0:
                rows_data.append((c0, c1, ""))
        return rows_data

    t1_rows = get_table_rows(1)
    t2_rows = get_table_rows(2)
    t3_rows = get_table_rows(3)

    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{form_info['panel_member']} - Approval Form</title>
<style>
  @page {{
    size: letter portrait;
    margin: 0.75in 0.85in;
    @bottom-center {{
      content: counter(page);
    }}
  }}
  body {{
    font-family: Arial, "Helvetica Neue", Helvetica, sans-serif;
    font-size: 10pt;
    line-height: 1.35;
    color: #000000;
    margin: 0;
    padding: 0;
  }}
  .header-top {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    font-size: 10pt;
    margin-bottom: 12pt;
  }}
  .school-title {{
    font-size: 10pt;
    font-weight: normal;
  }}
  .userinx-title {{
    font-size: 10pt;
    font-weight: bold;
    text-align: right;
  }}
  .form-title {{
    text-align: center;
    font-weight: bold;
    font-size: 11pt;
    margin: 0 0 2pt 0;
  }}
  .form-subtitle {{
    text-align: center;
    font-size: 9.5pt;
    font-style: italic;
    margin: 0 0 14pt 0;
  }}
  .greeting {{
    text-align: center;
    font-weight: bold;
    margin-bottom: 8pt;
  }}
  p.body-text {{
    text-align: justify;
    margin: 0 0 6pt 0;
    font-size: 9.5pt;
    line-height: 1.35;
  }}
  .rec-heading {{
    font-weight: bold;
    margin-top: 8pt;
    margin-bottom: 4pt;
    font-size: 9.5pt;
  }}
  .rec-list {{
    margin: 0 0 10pt 0;
    padding-left: 18pt;
    font-size: 9pt;
    line-height: 1.35;
  }}
  .rec-list li {{
    margin-bottom: 3pt;
  }}
  .rec-list strong {{
    font-weight: bold;
  }}
  table.form-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 8pt 0 14pt 0;
    font-size: 9pt;
  }}
  table.form-table th, table.form-table td {{
    border: 1px solid #000;
    padding: 5pt 7pt;
    vertical-align: top;
    line-height: 1.3;
  }}
  table.form-table th {{
    background-color: #f2f2f2;
    font-weight: bold;
    text-align: left;
  }}
  .section-label {{
    font-weight: bold;
    font-size: 10pt;
    margin-top: 12pt;
    margin-bottom: 4pt;
  }}
  .signature-block {{
    margin-top: 24pt;
    font-size: 9.5pt;
    line-height: 2.0;
    page-break-inside: avoid;
  }}
</style>
</head>
<body>

  <div class="header-top">
    <div class="school-title">School of Engineering</div>
    <div class="userinx-title">USERINX</div>
  </div>

  <div class="form-title">User Interface / User Experience Proposed Software Design Project Title Pre-approval Form</div>
  <div class="form-subtitle">(Refer to attached initial title proposal)</div>

  <div class="greeting">Dear Panel Member,</div>

  <p class="body-text">
    Attached is the Chapter 1 to 3 of the project proposal for USERINX CpE251 containing the approved Chapters 1 and 2, and Chapter 3: Information Architecture, Navigation Structure, User Flow, Wireframes, Design Standards, High Fidelity User Interface Design, and Usability Considerations.
  </p>
  <p class="body-text">
    Please determine if the attached project proposal is approved for proposal defense in terms of feasibility and relevance to CpE's field of specialization.
  </p>
  <p class="body-text">
    Once approved, the proponents are expected to proceed with system design documentation (Chapters 1 to 6) as a requirement for Database Management course.
  </p>

  <div class="rec-heading">Possible Recommendations:</div>
  <ul class="rec-list">
    <li><strong>Approved</strong> – The proposed title is accepted without revisions. The project team may proceed to the Proposal Defense.</li>
    <li><strong>Approved with Revisions</strong> – The proposed title is accepted subject to the completion of the required revisions. The project team may proceed to the Proposal Defense upon compliance with the panel's recommendations.</li>
    <li><strong>Change Topic</strong> – The proposed title is not approved. The project team shall propose a new project topic that addresses the recommendations of the evaluation panel.</li>
  </ul>
  <p class="body-text">
    If the project receives a Change Topic result during the scheduled Title Defense, the remaining defense schedule allocated for the UI/UX course shall be used for the Proposal Defense of the revised project. Any additional proposal defense, if required, shall be conducted during the next offering of the UI/UX course.
  </p>

  <!-- TABLE 0 -->
  <table class="form-table" style="margin-bottom: 12pt;">
    <thead>
      <tr>
        <th style="width: 72%;">PROPOSED TITLE</th>
        <th style="width: 28%; text-align: center;">RECOMMENDATION</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td style="font-weight: bold; padding: 8pt 7pt;">Development of a Centralized Curriculum Management and Mapping System with Role-Based Access Control and Audit Logging</td>
        <td></td>
      </tr>
    </tbody>
  </table>

  <!-- TABLE 1 -->
  <div class="section-label">Comments</div>
  <table class="form-table">
    <thead>
      <tr>
        <th style="width: 44%;">Comments</th>
        <th style="width: 44%;">Actions Taken</th>
        <th style="width: 12%; text-align: center;">Remarks</th>
      </tr>
    </thead>
    <tbody>
"""
    for c0, c1, c2 in t1_rows:
        html += f"""      <tr>
        <td>{c0.replace(chr(10), '<br>')}</td>
        <td>{c1}</td>
        <td></td>
      </tr>\n"""
        
    html += """    </tbody>
  </table>
"""

    if t2_rows:
        html += """  <!-- TABLE 2 -->
  <div class="section-label">For revisions</div>
  <table class="form-table">
    <thead>
      <tr>
        <th style="width: 44%;">Revisions</th>
        <th style="width: 44%;">Actions Taken</th>
        <th style="width: 12%; text-align: center;">Remarks</th>
      </tr>
    </thead>
    <tbody>
"""
        for c0, c1, c2 in t2_rows:
            html += f"""      <tr>
        <td>{c0.replace(chr(10), '<br>')}</td>
        <td>{c1}</td>
        <td></td>
      </tr>\n"""
        html += """    </tbody>
  </table>
"""

    if t3_rows:
        html += """  <!-- TABLE 3 -->
  <div class="section-label">Additional requirements (additional objectives that needs to be included in the complete proposal)</div>
  <table class="form-table">
    <thead>
      <tr>
        <th style="width: 44%;">Additional Requirements</th>
        <th style="width: 44%;">Actions Taken</th>
        <th style="width: 12%; text-align: center;">Remarks</th>
      </tr>
    </thead>
    <tbody>
"""
        for c0, c1, c2 in t3_rows:
            html += f"""      <tr>
        <td>{c0.replace(chr(10), '<br>')}</td>
        <td>{c1}</td>
        <td></td>
      </tr>\n"""
        html += """    </tbody>
  </table>
"""

    html += f"""  <div class="signature-block">
    <div><strong>NAME OF PANEL MEMBER:</strong> {form_info['panel_member']}&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;<strong>SIGNATURE :</strong> _________________________________</div>
    <div><strong>Date Signed :</strong> _________________________________</div>
  </div>

</body>
</html>
"""

    html_path = os.path.join(OUT_DIR, form_info["out_pdf_name"].replace(".pdf", ".html"))
    pdf_path = os.path.join(OUT_DIR, form_info["out_pdf_name"])
    
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
        
    url = "file:///" + os.path.abspath(html_path).replace("\\", "/")
    cmd = [
        EDGE_PATH,
        "--headless",
        "--disable-gpu",
        f"--user-data-dir={USER_DIR}",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        url
    ]
    subprocess.run(cmd, capture_output=True, text=True)
    if os.path.exists(html_path):
        os.remove(html_path)
    print(f"Created PDF: {pdf_path} (size: {os.path.getsize(pdf_path)})")

if __name__ == "__main__":
    for f_info in FORMS:
        recreate_docx(f_info)
        recreate_pdf(f_info)
    print("Recreation completed successfully.")
