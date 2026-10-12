import docx
import re

doc = docx.Document(r'C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Master_Cleaned.docx')
for i, p in enumerate(doc.paragraphs):
    if 'Figure 3.' in p.text:
        imgs = []
        for offset in [-1, 1]:
            idx = i + offset
            if 0 <= idx < len(doc.paragraphs):
                p_cand = doc.paragraphs[idx]
                for r in p_cand.runs:
                    if 'drawing' in r._r.xml:
                        blips = re.findall(r'r:embed="(rId\d+)"', r._r.xml)
                        for b in blips:
                            imgs.append((offset, b, doc.part.related_parts[b].partname))
        print(f"P[{i}] {p.text.strip()[:60]} -> {imgs}")
