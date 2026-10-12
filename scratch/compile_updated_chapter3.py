import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import subprocess

DOCX_SRC = r"C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Master_Cleaned.docx"
DOCX_OUT = r"C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Master_Updated.docx"
PDF_OUT = r"C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Master_Updated.pdf"
HTML_OUT = r"C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Master_Updated.html"

SCREEN_DIR = r"C:\Users\Dean Alejandrino\Downloads\KapTau_Chapter_3_Extracted\high_fidelity_screens_clean"
SCRATCH_DIR = r"c:\Users\Dean Alejandrino\Documents\antigravity\jolly-davinci\scratch"

# Ordered 24 screens with proper sequential numbering and descriptive text
ORDERED_SCREENS = [
    {
        "num": 37,
        "file": "01_login_modal.png",
        "title": "Authentication and Role-Based Login Screen",
        "desc": "The Authentication and Role-Based Login screen presents the initial institutional entry point for Asia Pacific College faculty and administrators. It features a custom native institutional credential authentication portal (institutional email and password), client-side validation, and direct administrative portal entry mechanisms."
    },
    {
        "num": 38,
        "file": "02_institutional_dashboard.png",
        "title": "Institutional Dashboard Screen",
        "desc": "The Institutional Dashboard screen provides an executive summary of academic metrics across Asia Pacific College. It displays high-level program statistics, quick-navigation cards for all academic schools, and system-wide curriculum status indicators."
    },
    {
        "num": 39,
        "file": "03_school_overview.png",
        "title": "School Overview Screen",
        "desc": "The School Overview screen organizes degree programs under the School of Engineering. It details executive leadership credentials, constituent undergraduate programs (Computer Engineering, Civil Engineering, Electronics Engineering), and institutional accreditation indicators."
    },
    {
        "num": 40,
        "file": "04_program_management.png",
        "title": "Program Management Directory Screen",
        "desc": "The Degree Program Management screen provides a catalog of engineering programs, student counts, and active curriculum cohorts. It acts as the operational bridge between school administration and individual degree curriculum management."
    },
    {
        "num": 41,
        "file": "05_curriculum_workbench.png",
        "title": "Curriculum Workbench Overview Screen",
        "desc": "The Curriculum Workbench Overview screen serves as the primary authoring workspace for the active Computer Engineering curriculum (CPE2026 / AY 2025-2026). It integrates version summary metrics, quick navigation into analytical subsystems, and export triggers."
    },
    {
        "num": 42,
        "file": "06_curriculum_flowchart.png",
        "title": "Interactive Curriculum Flowchart Screen",
        "desc": "The Interactive Curriculum Flowchart screen visualizes the four-year degree program across 12 trimesters as an interactive directed acyclic graph. It dynamically traces prerequisite trees, highlighting upstream feeder courses and downstream unlocked subjects upon course selection."
    },
    {
        "num": 43,
        "file": "07_course_details_modal.png",
        "title": "Course Details Modal Screen",
        "desc": "The Course Details Modal and Inspector screen provides a focused dialog detailing complete course metadata, including course code, descriptive title, lecture/laboratory credit unit split, hard prerequisites, co-requisites, and mapped student outcomes."
    },
    {
        "num": 44,
        "file": "08_curriculum_spreadsheet.png",
        "title": "Curriculum Spreadsheet Grid Screen",
        "desc": "The Curriculum Spreadsheet Data Grid screen delivers a high-density, Excel-style authoring table. It features an integrated formula bar, sticky column headers, term filters, real-time credit summation, and prerequisite code editing."
    },
    {
        "num": 45,
        "file": "09_obe_overview.png",
        "title": "OBE Hierarchy Hub Screen",
        "desc": "The Outcomes-Based Education (OBE) Hierarchy Hub screen organizes institutional educational alignment across four distinct tiers: Graduate Attributes (GA), Program Educational Objectives (PEO), Student Outcomes (SO), and Course Competencies."
    },
    {
        "num": 46,
        "file": "10_obe_course_so.png",
        "title": "OBE Course-to-SO Matrix Screen",
        "desc": "The Course-to-Student Outcomes (I-E-D) Matrix screen renders the comprehensive mapping table aligning all degree courses against the 13 CHED CMO No. 92 Student Outcomes, tagging each course with Introductory (I), Enabling (E), or Demonstrative (D) levels."
    },
    {
        "num": 47,
        "file": "11_obe_so_peo.png",
        "title": "OBE SO-to-PEO Alignment Schedule Screen",
        "desc": "The Student Outcomes to Program Educational Objectives (SO-to-PEO) Alignment Schedule maps the 13 Student Outcomes against the 5 Program Educational Objectives, ensuring rigorous alignment between classroom performance and career competencies."
    },
    {
        "num": 48,
        "file": "12_obe_peo_ga.png",
        "title": "OBE PEO-to-GA Alignment Screen",
        "desc": "The PEO to Graduate Attributes (GA) Alignment screen connects program-level educational objectives directly to Asia Pacific College's institutional graduate attributes, validating that program learning goals realize institutional mission and values."
    },
    {
        "num": 49,
        "file": "14_official_curriculum_flowchart.png",
        "title": "Official Curriculum Flowchart Document Screen",
        "desc": "The Official Document 1: Curriculum Flowchart Document screen formats the complete curriculum dependency matrix for standardized registrar paper publication, featuring official APC institutional headers and term grouping bands."
    },
    {
        "num": 50,
        "file": "17_official_prospectus.png",
        "title": "Official Prospectus Sheet Screen",
        "desc": "The Official Document 2: Student Prospectus Sheet screen delivers the authoritative student curriculum checklist, displaying academic requirements, term credit counts, grade tracking slots, and student number audit columns."
    },
    {
        "num": 51,
        "file": "13_official_course_catalog.png",
        "title": "Official Course Catalog Document Screen",
        "desc": "The Official Document 3: Course Catalog screen provides a print-ready directory of all courses within the academic program, cataloging official course codes, complete descriptive titles, credit unit distributions, and explicit requisite requirements."
    },
    {
        "num": 52,
        "file": "16_official_program_of_study.png",
        "title": "Official Program of Study Sheet Screen",
        "desc": "The Official Document 4: Program of Study Sheet screen formats the academic sequence into a concise registrar schedule, showing prescribed term enrollment loads and total cumulative degree credit allocations."
    },
    {
        "num": 53,
        "file": "15_official_comparative_summary.png",
        "title": "Official Comparative Summary Report Screen",
        "desc": "The Official Document 5: Comparative Summary Report screen presents an analytical comparison between curriculum versions, summarizing changes in credit unit allocations, prerequisite changes, and course modifications across revisions."
    },
    {
        "num": 54,
        "file": "24_version_history.png",
        "title": "Curriculum Version History and Revision Log Screen",
        "desc": "The Curriculum Version History and Revision Log screen tracks institutional curriculum releases across academic cohorts (e.g., CPE2026, CPE2027 draft), providing version freezing, revision diffing, and baseline restoration controls."
    },
    {
        "num": 55,
        "file": "20_task_delegation.png",
        "title": "Collaborative Task Delegation Console Screen",
        "desc": "The Collaborative Task Delegation Console screen empowers Program Directors to assign time-bound editing permissions to faculty leads across specialized academic clusters, complete with expiration limits and task status monitoring."
    },
    {
        "num": 56,
        "file": "21_system_audit_trail.png",
        "title": "Immutable System Audit Trail Screen",
        "desc": "The Immutable System Audit Trail screen provides a tamper-evident chronological ledger of all user actions, curriculum edits, approval state transitions, and system events with precise UTC timestamps and actor identification."
    },
    {
        "num": 57,
        "file": "18_batch_import_modal.png",
        "title": "Batch Excel Data Ingestion Modal Screen",
        "desc": "The Batch Excel Data Ingestion Modal screen provides an automated pipeline for ingesting official registrar spreadsheet files (.xlsx), performing instant schema validation, and populating entire degree programs in a single operation."
    },
    {
        "num": 58,
        "file": "19_data_export_backup_modal.png",
        "title": "Data Export and Institutional Backup Modal Screen",
        "desc": "The Data Export and Institutional Backup Modal screen facilitates institutional data preservation, enabling one-click exports of curriculum datasets in structured JSON, Excel (.xlsx), and PDF document formats."
    },
    {
        "num": 59,
        "file": "22_system_settings_modal.png",
        "title": "System Preferences and Settings Modal Screen",
        "desc": "The System Preferences and Settings Modal screen provides administrative controls for theme toggling (Light / Obsidian Dark Mode), user role simulation, session persistence management, and system diagnostic testing."
    },
    {
        "num": 60,
        "file": "23_dark_mode_flowchart.png",
        "title": "Obsidian Gold Dark Mode Flowchart Screen",
        "desc": "The Obsidian Gold High-Contrast Dark Mode Screen renders the complete curriculum matrix against deep slate surfaces (#0B1120) with luminous gold accents and glowing prerequisite tracer pathways, reducing eye fatigue during prolonged analytical sessions."
    }
]

print("Loading base document...")
doc = docx.Document(DOCX_SRC)

# Locate Section 3.6 (P[221] to P[294])
p_start = None
p_end = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "High Fidelity User Interface Design":
        p_start = i
    elif p.text.strip() == "Usability Considerations":
        p_end = i
        break

print(f"Section 3.6 spans from P[{p_start}] to P[{p_end}]")

# We will remove the old 3.6 paragraphs (p_start + 2 to p_end - 1) and insert the new ordered content
# p_start is the Heading 3
# p_start + 1 is the introductory paragraph
intro_p = doc.paragraphs[p_start + 1]

# Delete paragraphs between p_start + 2 and p_end
for i in range(p_end - 1, p_start + 1, -1):
    p = doc.paragraphs[i]
    p._element.getparent().remove(p._element)

print("Removed old 3.6 figure paragraphs.")

# Now insert each ordered screen right after intro_p
# In python-docx, insert_paragraph_before() on usability heading inserts in front of Usability Considerations!
usability_p = doc.paragraphs[p_start + 2] # This is now Usability Considerations!

for screen in ORDERED_SCREENS:
    # 1. Image paragraph
    img_p = usability_p.insert_paragraph_before()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_path = os.path.join(SCREEN_DIR, screen["file"])
    if os.path.exists(img_path):
        run = img_p.add_run()
        run.add_picture(img_path, width=Inches(6.5))
    
    # 2. Caption paragraph
    cap_p = usability_p.insert_paragraph_before()
    cap_p.style = 'Caption'
    cap_p.paragraph_format.space_before = Pt(4)
    cap_p.paragraph_format.space_after = Pt(4)
    cap_run = cap_p.add_run(f"Figure 3.{screen['num']}. {screen['title']}")
    cap_run.bold = True
    
    # 3. Description paragraph
    desc_p = usability_p.insert_paragraph_before()
    desc_p.style = 'Normal'
    desc_p.paragraph_format.space_before = Pt(0)
    desc_p.paragraph_format.space_after = Pt(12)
    desc_run_label = desc_p.add_run(f"{screen['title']}: ")
    desc_run_label.bold = True
    desc_run_text = desc_p.add_run(screen["desc"])

print("Successfully inserted all 24 ordered screens into document.")
doc.save(DOCX_OUT)
print(f"Saved updated DOCX to: {DOCX_OUT}")

# Convert DOCX to HTML using mammoth
import mammoth
import re

with open(DOCX_OUT, "rb") as docx_file:
    result = mammoth.convert_to_html(docx_file)
    html_body = result.value

# Style captions and tables
html_body = re.sub(
    r'<p>(Figure 3\.\d+\..*?)</p>',
    r'<p style="text-align: center; font-weight: bold; margin-top: 6pt; margin-bottom: 12pt; font-size: 11pt;">\1</p>',
    html_body
)
html_body = re.sub(
    r'<p>(Table 3\.\d+.*?)</p>',
    r'<p style="font-weight: bold; margin-top: 14pt; margin-bottom: 4pt; font-size: 11pt;">\1</p>',
    html_body
)

html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>KapTau - Chapter 3</title>
<style>
  @page {{
    size: letter portrait;
    margin: 1.0in;
  }}
  body {{
    font-family: 'Times New Roman', Times, serif;
    font-size: 12pt;
    line-height: 1.5;
    color: #111;
    margin: 0;
    padding: 0;
  }}
  h2 {{
    font-size: 15pt;
    font-weight: bold;
    margin-top: 24pt;
    margin-bottom: 8pt;
    page-break-after: avoid;
    break-after: avoid;
  }}
  h3 {{
    font-size: 13.5pt;
    font-weight: bold;
    margin-top: 18pt;
    margin-bottom: 6pt;
    page-break-after: avoid;
    break-after: avoid;
  }}
  p {{
    margin-top: 0;
    margin-bottom: 8pt;
    text-align: justify;
  }}
  img {{
    max-width: 100%;
    height: auto;
    display: block;
    margin: 12pt auto 4pt auto;
    border: 1px solid #cbd5e1;
    border-radius: 4px;
    page-break-inside: avoid;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    margin: 10pt 0;
    page-break-inside: avoid;
    font-size: 10.5pt;
  }}
  th, td {{
    border: 1px solid #333;
    padding: 6pt 8pt;
    text-align: left;
    vertical-align: top;
  }}
  th {{
    background-color: #f1f5f9;
    font-weight: bold;
  }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""

with open(HTML_OUT, "w", encoding="utf-8") as f:
    f.write(html_template)
print(f"Saved rendered HTML to: {HTML_OUT}")

# Render HTML to PDF using Edge headless
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
if not os.path.exists(EDGE_PATH):
    EDGE_PATH = r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"

edge_cmd = [
    EDGE_PATH,
    "--headless",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={PDF_OUT}",
    f"file:///{os.path.abspath(HTML_OUT).replace(os.sep, '/')}"
]
edge_res = subprocess.run(edge_cmd, capture_output=True, text=True)
if edge_res.returncode == 0:
    print(f"Generated PDF successfully at {PDF_OUT}.")
    print("PDF size:", os.path.getsize(PDF_OUT), "bytes")
else:
    print("Edge PDF error:", edge_res.stderr)
